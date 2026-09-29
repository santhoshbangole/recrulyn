"""Recrulyn go-to-market pitch deck (16:9)."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import nsmap, qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
LOGO = ROOT / "public" / "Logo-Monogram.png"
OUT = Path(__file__).resolve().parent / "Recrulyn_Go_To_Market_Pitch.pptx"

# Brand
INK = RGBColor(0x11, 0x18, 0x27)
INK2 = RGBColor(0x37, 0x41, 0x51)
MUTED = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = RGBColor(0xFB, 0xF9, 0xF4)
ORANGE = RGBColor(0xF2, 0x5C, 0x05)
ORANGE_DK = RGBColor(0xC4, 0x45, 0x00)
GREEN = RGBColor(0x2F, 0x7D, 0x4A)
GREEN_DK = RGBColor(0x1F, 0x5C, 0x36)
NAVY = RGBColor(0x0E, 0x15, 0x26)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xE5, 0xE7, 0xEB)
TEAL = RGBColor(0x0F, 0x8B, 0x8D)

W, H = Inches(13.333), Inches(7.5)


def set_run(run, text, size=18, bold=False, color=INK, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_text(shape, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Calibri"):
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    set_run(p.add_run() if p.runs else p.runs[0] if False else _ensure_run(p), text, size, bold, color, font)
    return tf


def _ensure_run(p):
    if p.runs:
        return p.runs[0]
    return p.add_run()


def box(slide, l, t, w, h, fill=None, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    s.line.fill.background()
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    return s


def round_box(slide, l, t, w, h, fill=WHITE, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    s.adjustments[0] = 0.08
    s.line.fill.background()
    if line is not None:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    return s


def tb(slide, l, t, w, h, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Calibri"):
    s = slide.shapes.add_textbox(l, t, w, h)
    tf = s.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    set_run(r, text, size, bold, color, font)
    return s


def multi(slide, l, t, w, h, lines, size=15, color=INK2, bold=False, space=8):
    s = slide.shapes.add_textbox(l, t, w, h)
    tf = s.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(space)
        r = p.add_run()
        set_run(r, line, size, bold, color)
    return s


def footer(slide, n, total=16):
    box(slide, 0, Inches(7.28), W, Inches(0.22), NAVY)
    tb(slide, Inches(0.4), Inches(7.28), Inches(8), Inches(0.22), "RECRULYN  ·  Confidential pitch", 10, False, WHITE)
    tb(slide, Inches(11.4), Inches(7.28), Inches(1.6), Inches(0.22), f"{n}  /  {total}", 10, False, WHITE, PP_ALIGN.RIGHT)


def header_bar(slide, kicker, title, subtitle=None):
    box(slide, 0, 0, W, Inches(1.28), NAVY)
    box(slide, 0, 0, Inches(0.12), Inches(1.28), ORANGE)
    tb(slide, Inches(0.45), Inches(0.12), Inches(12), Inches(0.28), kicker.upper(), 11, True, ORANGE, font="Calibri")
    tb(slide, Inches(0.45), Inches(0.38), Inches(12.4), Inches(0.48), title, 26, True, WHITE)
    if subtitle:
        tb(slide, Inches(0.45), Inches(0.88), Inches(12.4), Inches(0.32), subtitle, 13, False, RGBColor(0xCB, 0xD5, 0xE1))


def card_title(slide, shape, title, sub, y_off=0.12):
    l, t, w = shape.left, shape.top, shape.width
    tb(slide, l + Inches(0.18), t + Inches(y_off), w - Inches(0.36), Inches(0.32), title, 15, True, INK)
    if sub:
        tb(slide, l + Inches(0.18), t + Inches(y_off + 0.32), w - Inches(0.36), Inches(1.4), sub, 12, False, INK2)


def bg(slide, color=CREAM):
    box(slide, 0, 0, W, H, color)


def add_logo(slide, l, t, w=Inches(0.72)):
    if LOGO.exists():
        slide.shapes.add_picture(str(LOGO), l, t, w, w)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]
    total = 16

    # ── 1 Title ──────────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    bg(s, NAVY)
    box(s, 0, 0, Inches(0.18), H, ORANGE)
    add_logo(s, Inches(0.55), Inches(0.45), Inches(0.9))
    tb(s, Inches(1.6), Inches(0.62), Inches(8), Inches(0.5), "RECRULYN", 22, True, WHITE)
    tb(s, Inches(0.55), Inches(1.7), Inches(12), Inches(1.4),
       "The AI HR workspace that turns incoming mail into a ready talent pipeline.",
       32, True, WHITE)
    tb(s, Inches(0.55), Inches(3.35), Inches(11.2), Inches(1.1),
       "Pitch for universities, institutes, and industry partners.\nHire faster. Paperwork in one click. HR never hunts an inbox again.",
       18, False, RGBColor(0xCB, 0xD5, 0xE1))

    pills = [
        (Inches(0.55), "IMAP → Candidate"),
        (Inches(3.35), "AI shortlist"),
        (Inches(5.7), "LOA / NDA packs"),
        (Inches(8.2), "HR inbox, live"),
    ]
    for x, label in pills:
        sh = round_box(s, x, Inches(4.75), Inches(2.55), Inches(0.48), ORANGE)
        tb(s, x, Inches(4.82), Inches(2.55), Inches(0.38), label, 13, True, WHITE, PP_ALIGN.CENTER)

    tb(s, Inches(0.55), Inches(6.55), Inches(10), Inches(0.4),
       "Go-to-market  ·  Confidential  ·  2026", 13, False, MUTED)
    box(s, 0, Inches(7.28), W, Inches(0.22), ORANGE)

    # ── 2 Agenda ─────────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    bg(s)
    header_bar(s, "Agenda", "What this conversation covers")
    items = [
        ("01", "The HR time trap", "Why campus and industry hiring still burns days in mail, PDFs, and copy-paste."),
        ("02", "Product, live today", "Features already enabled in Recrulyn — pipeline, AI recruiter, documents, leave, analytics."),
        ("03", "IMAP mail engine", "How the HR mailbox becomes structured candidate records without manual filing."),
        ("04", "Go-to-market", "How we approach universities, institutes, and industry in parallel."),
        ("05", "Partnership ask", "Pilot, campus license, and industry deployment models."),
    ]
    for i, (num, title, desc) in enumerate(items):
        y = Inches(1.52) + Inches(i * 1.05)
        c = round_box(s, Inches(0.45), y, Inches(12.4), Inches(0.92), WHITE, LINE)
        nbox = round_box(s, Inches(0.62), y + Inches(0.2), Inches(0.72), Inches(0.52), ORANGE)
        tb(s, Inches(0.62), y + Inches(0.28), Inches(0.72), Inches(0.4), num, 16, True, WHITE, PP_ALIGN.CENTER)
        tb(s, Inches(1.55), y + Inches(0.14), Inches(10.8), Inches(0.36), title, 18, True, INK)
        tb(s, Inches(1.55), y + Inches(0.48), Inches(10.8), Inches(0.36), desc, 13, False, INK2)
    footer(s, 2, total)

    # ── 3 Problem ────────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    bg(s)
    header_bar(s, "The problem", "HR still does the work that software should do",
               "Especially on campus drives, intern batches, and high-volume industry hiring.")
    pains = [
        ("Inbox as ATS", "Resumes arrive as attachments across INBOX, HR folders, and forwarded threads. Nothing is a candidate until someone opens it."),
        ("Copy-paste profiles", "Name, email, skills, education, and experience are typed by hand into sheets or another tool."),
        ("JD matching is tribal", "Shortlisting depends on who has time to read 80 CVs before tomorrow’s interview panel."),
        ("Paper after the yes", "LOA, NDA, LOC, LOR, certificates, and internship letters live in Word templates and WhatsApp."),
        ("Leaders wait", "Management cannot see pipeline, leave, or hiring risk without asking HR for a screenshot."),
        ("Campus + industry dual load", "Universities send batches; companies send JDs. HR is the bottleneck in both directions."),
    ]
    for i, (title, desc) in enumerate(pains):
        col, row = i % 3, i // 3
        x, y = Inches(0.4) + Inches(col * 4.25), Inches(1.55) + Inches(row * 2.55)
        c = round_box(s, x, y, Inches(4.05), Inches(2.35), WHITE, LINE)
        box(s, x, y, Inches(0.1), Inches(2.35), ORANGE)
        tb(s, x + Inches(0.28), y + Inches(0.22), Inches(3.55), Inches(0.5), title, 16, True, INK)
        tb(s, x + Inches(0.28), y + Inches(0.78), Inches(3.55), Inches(1.35), desc, 13, False, INK2)
    footer(s, 3, total)

    # ── 4 Solution ───────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    bg(s)
    header_bar(s, "The idea", "Recrulyn is the HR operating system for hiring, people, and paperwork")
    tb(s, Inches(0.5), Inches(1.5), Inches(12.3), Inches(0.7),
       "One workspace instead of six tools: mail, ATS, AI scoring, document packs, leave, and leadership briefings.",
       16, False, INK2)

    pillars = [
        (GREEN, "Capture", "IMAP syncs the HR mailbox. Resumes and HR mail become records automatically."),
        (TEAL, "Understand", "Parsers + AI extract skills, education, experience and score against the job."),
        (ORANGE, "Act", "HR reviews, replies from templates, issues LOA/NDA, and briefs management."),
        (NAVY, "Scale", "Same engine for campus internships, institute placements, and industry roles."),
    ]
    for i, (color, title, desc) in enumerate(pillars):
        x = Inches(0.4) + Inches(i * 3.2)
        c = round_box(s, x, Inches(2.35), Inches(3.0), Inches(3.55), WHITE, LINE)
        box(s, x, Inches(2.35), Inches(3.0), Inches(0.12), color)
        tb(s, x + Inches(0.18), Inches(2.65), Inches(2.64), Inches(0.4), f"0{i+1}", 14, True, color)
        tb(s, x + Inches(0.18), Inches(3.1), Inches(2.64), Inches(0.55), title, 22, True, INK)
        tb(s, x + Inches(0.18), Inches(3.75), Inches(2.64), Inches(1.8), desc, 14, False, INK2)
    footer(s, 4, total)

    # ── 5 Features ───────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    bg(s)
    header_bar(s, "Product — enabled today", "Features HR can use without waiting for a roadmap")
    feats = [
        ("Talent pipeline", "Candidate records, status, activity log, departments, and open roles."),
        ("HR inbox", "Live IMAP import of mail + attachments. Open an email and hydrate body + files on demand."),
        ("Resume AI", "Parse PDF/DOCX: name, contact, skills, education, projects, certifications."),
        ("AI recruiter", "Natural-language role + skills → ranked shortlist against stored candidates."),
        ("Resume intelligence", "Score a CV against a job description with match history."),
        ("Issue letters", "One-click LOA, NDA, LOC, LOR, certificates — with approvals."),
        ("Email templates", "Interview invite, internship confirmation, and outbound SMTP from HR."),
        ("People ops", "Leave desk, multi-level approvals, employee records, role-based access."),
        ("People analytics", "Dashboards for HR, management, and employees — hiring and leave in one view."),
        ("Document vault", "Store resumes and generated packs (local + WorkDrive)."),
        ("Reudey copilot", "In-app assistant for recruiter questions and workflow help."),
        ("Secure roles", "HR, Employee, Management, Admin — each sees only what they should."),
    ]
    for i, (title, desc) in enumerate(feats):
        col, row = i % 4, i // 4
        x, y = Inches(0.32) + Inches(col * 3.24), Inches(1.48) + Inches(row * 1.85)
        c = round_box(s, x, y, Inches(3.1), Inches(1.72), WHITE, LINE)
        tb(s, x + Inches(0.16), y + Inches(0.16), Inches(2.78), Inches(0.5), title, 14, True, GREEN_DK)
        tb(s, x + Inches(0.16), y + Inches(0.64), Inches(2.78), Inches(0.95), desc, 12, False, INK2)
    footer(s, 5, total)

    # ── 6 IMAP ───────────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    bg(s)
    header_bar(s, "How mail becomes hiring data", "IMAP is the silent intake layer — HR never downloads a zip of CVs")
    steps = [
        ("1. Connect", "Secure IMAP (port 993) to the HR mailbox — e.g. Zoho. SMTP for outbound."),
        ("2. Sync folders", "INBOX + HR-labelled folders. Last ~21 days. Deduped by mailbox + UID."),
        ("3. Parse mail", "Subject, sender, body, date. Skip logos/tracking pixels."),
        ("4. Classify files", "PDF / DOCX / ID photos → resume vs attachment. Store in vault."),
        ("5. Extract", "Resume text → name, email, phone, skills, education, experience."),
        ("6. HR sees it", "Appears in HR Inbox + Talent pipeline. Click to open full body from IMAP."),
    ]
    for i, (title, desc) in enumerate(steps):
        col, row = i % 3, i // 3
        x, y = Inches(0.4) + Inches(col * 4.25), Inches(1.5) + Inches(row * 2.15)
        c = round_box(s, x, y, Inches(4.05), Inches(1.95), WHITE, LINE)
        bar = box(s, x, y, Inches(4.05), Inches(0.08), ORANGE if row == 0 else GREEN)
        tb(s, x + Inches(0.2), y + Inches(0.22), Inches(3.65), Inches(0.4), title, 16, True, INK)
        tb(s, x + Inches(0.2), y + Inches(0.68), Inches(3.65), Inches(1.1), desc, 13, False, INK2)
    tb(s, Inches(0.45), Inches(5.95), Inches(12.4), Inches(1.05),
       "On open: Recrulyn can re-fetch the original message by mailbox + IMAP UID (or match sender/subject/time) so HR always has the real attachment — not a stale copy. Processed UIDs are marked so the same mail is never imported twice.",
       14, False, INK2)
    footer(s, 6, total)

    # ── 7 IMAP diagram ───────────────────────────────────────
    s = prs.slides.add_slide(blank)
    bg(s)
    header_bar(s, "Data path", "From campus / candidate mail to the HR dashboard")
    nodes = [
        (0.35, "Candidate / TPO\nsends CV to HR mail"),
        (2.85, "IMAP sync\nHR mailbox"),
        (5.35, "Parse + vault\nresume + IDs"),
        (7.85, "AI extract\nprofile fields"),
        (10.35, "HR workspace\npipeline + inbox"),
    ]
    for i, (x, label) in enumerate(nodes):
        sh = round_box(s, Inches(x), Inches(1.7), Inches(2.35), Inches(1.45), NAVY if i == 4 else WHITE, LINE if i != 4 else None)
        if i == 4:
            sh.fill.solid()
            sh.fill.fore_color.rgb = GREEN
        tb(s, Inches(x), Inches(1.95), Inches(2.35), Inches(1.05), label, 13, True, WHITE if i == 4 else INK, PP_ALIGN.CENTER)
        if i < 4:
            arr = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 2.38), Inches(2.22), Inches(0.42), Inches(0.32))
            arr.fill.solid()
            arr.fill.fore_color.rgb = ORANGE
            arr.line.fill.background()

    # two columns below
    left = round_box(s, Inches(0.35), Inches(3.45), Inches(6.15), Inches(3.5), WHITE, LINE)
    right = round_box(s, Inches(6.7), Inches(3.45), Inches(6.2), Inches(3.5), WHITE, LINE)
    tb(s, Inches(0.55), Inches(3.6), Inches(5.8), Inches(0.4), "Without Recrulyn", 16, True, ORANGE)
    multi(s, Inches(0.55), Inches(4.1), Inches(5.8), Inches(2.6), [
        "• Open 40 emails. Download 40 PDFs.",
        "• Rename files. Drop into a shared drive.",
        "• Type names into Excel / another ATS.",
        "• Miss replies sitting in a subfolder.",
        "• Re-attach the same CV when issuing LOA.",
        "• Leadership asks “how many applied?” — wait.",
    ], 14, INK2, space=6)
    tb(s, Inches(6.9), Inches(3.6), Inches(5.8), Inches(0.4), "With Recrulyn IMAP + HR inbox", 16, True, GREEN)
    multi(s, Inches(6.9), Inches(4.1), Inches(5.8), Inches(2.6), [
        "• Mail lands → candidate record exists.",
        "• Resume text and URL already attached.",
        "• HR reviews in one inbox, not Outlook + Drive.",
        "• AI ranks against the open role / JD.",
        "• Reply from templates; issue LOA/NDA from the same profile.",
        "• Dashboard counts are live for HR and management.",
    ], 14, INK2, space=6)
    footer(s, 7, total)

    # ── 8 Direct HR access ───────────────────────────────────
    s = prs.slides.add_slide(blank)
    bg(s)
    header_bar(s, "Direct access for HR", "One login. The work is already waiting — not buried in mail.")
    rows = [
        ("HR Inbox", "See imported mail, sender, subject, attachments, resume link. Open to hydrate body from IMAP."),
        ("Talent pipeline", "Every extracted person is a candidate: status, notes, activity, department fit."),
        ("Open roles", "Job requirements become the scoring target for AI recruiter and resume intelligence."),
        ("AI recruiter", "Type the role and must-have skills. Get ranked people from what already arrived."),
        ("Issue letters", "Generate LOA / NDA / LOC / LOR / certificate from the candidate — no retyping."),
        ("Approvals & leave", "HR desk plus management sign-off. Employees self-serve leave without ticket chaos."),
    ]
    for i, (title, desc) in enumerate(rows):
        y = Inches(1.48) + Inches(i * 0.88)
        c = round_box(s, Inches(0.4), y, Inches(12.5), Inches(0.8), WHITE, LINE)
        box(s, Inches(0.4), y, Inches(0.12), Inches(0.8), GREEN if i % 2 == 0 else ORANGE)
        tb(s, Inches(0.75), y + Inches(0.08), Inches(2.6), Inches(0.64), title, 15, True, INK)
        tb(s, Inches(3.4), y + Inches(0.12), Inches(9.2), Inches(0.58), desc, 14, False, INK2)
    footer(s, 8, total)

    # ── 9 Time saved ─────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    bg(s)
    header_bar(s, "Time and effort", "Where hours disappear — and where Recrulyn gives them back")
    headers = ["HR task", "Typical today", "With Recrulyn", "HR effort"]
    xs = [0.35, 3.15, 6.55, 10.15]
    ws = [2.7, 3.3, 3.5, 2.75]
    for x, w, htxt in zip(xs, ws, headers):
        b = box(s, Inches(x), Inches(1.5), Inches(w), Inches(0.48), NAVY)
        tb(s, Inches(x), Inches(1.56), Inches(w), Inches(0.4), htxt, 13, True, WHITE, PP_ALIGN.CENTER)

    table = [
        ["Intake CVs from mail", "Download, rename, file", "IMAP auto-import + vault", "Minutes → seconds"],
        ["Build candidate card", "Read CV, type fields", "Parser fills profile", "~80% less typing"],
        ["Shortlist vs JD", "Read all CVs overnight", "AI score + ranked list", "Panel-ready same day"],
        ["Reply / interview mail", "Draft from scratch", "Templates + SMTP send", "Repeatable in clicks"],
        ["Offer / intern pack", "Edit Word, chase sign-off", "LOA/NDA + approval flow", "One workspace"],
        ["Status to leadership", "Build a slide from Excel", "Live dashboards", "Always current"],
    ]
    for r, row in enumerate(table):
        y = Inches(2.05) + Inches(r * 0.78)
        fill = WHITE if r % 2 == 0 else RGBColor(0xF3, 0xF4, 0xF6)
        for x, w, cell, j in zip(xs, ws, row, range(4)):
            box(s, Inches(x), y, Inches(w), Inches(0.78), fill)
            col = GREEN_DK if j == 3 else INK
            tb(s, Inches(x) + Inches(0.08), y + Inches(0.18), Inches(w) - Inches(0.12), Inches(0.48), cell, 12, j == 0, col, PP_ALIGN.LEFT)
    footer(s, 9, total)

    # ── 10 Universities ──────────────────────────────────────
    s = prs.slides.add_slide(blank)
    bg(s)
    header_bar(s, "Go-to-market — universities", "Placement cells send talent. Recrulyn receives it as structure, not clutter.")
    tb(s, Inches(0.45), Inches(1.48), Inches(12.4), Inches(0.55),
       "Approach: TPO / Placement Officer / Dean (Academics)  ·  Pilot one drive or one internship season.",
       14, False, INK2)
    offers = [
        ("For the university", [
            "Students apply to a published HR mail — no extra portal to learn on day one.",
            "Internship confirmation, NDA, and joining letters issued in the same system companies already use.",
            "Cleaner audit trail for MoUs: who applied, who interned, which letter went out.",
            "Placement reports without Excel archaeology after the drive.",
        ]),
        ("How we land campus", [
            "Start with internships and project interns (high volume, high paperwork).",
            "Co-brand a “campus intake week”: mailbox + JD + AI shortlist demo on real CVs.",
            "Offer a placement-cell login (read / analytics) plus HR login for the host company.",
            "Expand to full-time campus hiring once the intern pipeline is trusted.",
        ]),
        ("Proof they feel", [
            "A batch of 50 CVs in mail → 50 parsed profiles before the panel sits.",
            "Match to the intern JD (skills, education, projects).",
            "Same afternoon: interview invites from templates.",
            "Selected students: LOA + NDA pack without Word mail-merge.",
        ]),
    ]
    for i, (title, bullets) in enumerate(offers):
        x = Inches(0.35) + Inches(i * 4.3)
        c = round_box(s, x, Inches(2.15), Inches(4.1), Inches(4.7), WHITE, LINE)
        box(s, x, Inches(2.15), Inches(4.1), Inches(0.1), ORANGE)
        tb(s, x + Inches(0.2), Inches(2.4), Inches(3.7), Inches(0.45), title, 16, True, INK)
        multi(s, x + Inches(0.2), Inches(2.95), Inches(3.7), Inches(3.6), ["• " + b for b in bullets], 13, INK2, space=10)
    footer(s, 10, total)

    # ── 11 Institutes ────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    bg(s)
    header_bar(s, "Go-to-market — institutes & skilling", "Polytechnics, training academies, and bootcamps need employability, not another LMS.")
    left = round_box(s, Inches(0.4), Inches(1.5), Inches(6.2), Inches(5.4), WHITE, LINE)
    right = round_box(s, Inches(6.75), Inches(1.5), Inches(6.15), Inches(5.4), WHITE, LINE)
    tb(s, Inches(0.6), Inches(1.7), Inches(5.8), Inches(0.4), "Why institutes partner", 18, True, INK)
    multi(s, Inches(0.6), Inches(2.25), Inches(5.8), Inches(4.4), [
        "• Graduating cohorts have CVs of uneven quality — Recrulyn still extracts skills and projects.",
        "• Industry partners want a ranked list, not a ZIP file of 200 resumes.",
        "• Certificates and LORs can be issued from the same people record used for hiring.",
        "• Training-to-job handoff: institute shares a shortlist; company HR sees it in Recrulyn.",
        "• Position Recrulyn as the “last mile” of skilling: evidence of job-readiness.",
        "• Bundle: institute license for resume intelligence + partner-company HR inbox.",
    ], 14, INK2, space=8)
    tb(s, Inches(6.95), Inches(1.7), Inches(5.8), Inches(0.4), "Motion we propose", 18, True, INK)
    multi(s, Inches(6.95), Inches(2.25), Inches(5.75), Inches(4.4), [
        "1. Workshop: “Your mailbox is your ATS” — live IMAP demo.",
        "2. Ingest one graduating batch against 1–2 industry JDs.",
        "3. Joint shortlist review with a hiring manager in the room.",
        "4. Issue sample LOR / certificate / intern LOA.",
        "5. MoU: institute feeds talent; Recrulyn is the shared operating layer.",
        "6. Upsell industry tenants who hire from that institute every season.",
    ], 14, INK2, space=8)
    footer(s, 11, total)

    # ── 12 Industry ──────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    bg(s)
    header_bar(s, "Go-to-market — industry", "SMEs and mid-market teams that hire from mail, referrals, and campus — not from a giant ATS budget.")
    segs = [
        ("Who we sell to", "HR managers, founders who still do HR, TA leads in 20–500 people companies, shared-services HR for groups."),
        ("Entry wedge", "Connect existing HR mailbox. Value in week one: inbox is already a pipeline. No rip-and-replace."),
        ("Land", "IMAP + pipeline + AI shortlist on current open roles. Prove time saved on the next 25 applicants."),
        ("Expand", "Document packs (LOA/NDA), leave & approvals, management dashboards, multi-department roles."),
        ("Keep", "Every new mail makes the talent graph richer. Switching cost is the living inbox + document history."),
        ("Channels", "Campus MoUs, institute alumni hiring, industry associations, Zoho/mail-first IT teams, HR consultancies."),
    ]
    for i, (title, desc) in enumerate(segs):
        col, row = i % 3, i // 3
        x, y = Inches(0.4) + Inches(col * 4.25), Inches(1.52) + Inches(row * 2.55)
        c = round_box(s, x, y, Inches(4.05), Inches(2.35), WHITE, LINE)
        tb(s, x + Inches(0.22), y + Inches(0.22), Inches(3.6), Inches(0.45), title, 16, True, ORANGE if row == 0 else GREEN)
        tb(s, x + Inches(0.22), y + Inches(0.75), Inches(3.6), Inches(1.4), desc, 14, False, INK2)
    footer(s, 12, total)

    # ── 13 GTM motion ────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    bg(s)
    header_bar(s, "One motion, three doors", "Universities, institutes, and industry share the same engine")
    # funnel-ish 3 columns
    cols = [
        (NAVY, "University", "Door: TPO / Dean", [
            "Campus drive or intern season",
            "Shared HR mailbox for applications",
            "Placement analytics for the cell",
            "Outcome: hired intern + letters",
        ]),
        (ORANGE, "Institute", "Door: Training head", [
            "Cohort CVs vs industry JD",
            "Job-readiness shortlist",
            "LOR / certificate issuance",
            "Outcome: employer offtake",
        ]),
        (GREEN, "Industry", "Door: HR / founder", [
            "Connect IMAP in days",
            "AI recruiter on live applicants",
            "People docs + leave later",
            "Outcome: hours back every week",
        ]),
    ]
    for i, (color, title, door, bullets) in enumerate(cols):
        x = Inches(0.4) + Inches(i * 4.25)
        top = box(s, x, Inches(1.5), Inches(4.05), Inches(1.15), color)
        tb(s, x, Inches(1.58), Inches(4.05), Inches(0.45), title, 20, True, WHITE, PP_ALIGN.CENTER)
        tb(s, x, Inches(2.05), Inches(4.05), Inches(0.45), door, 13, False, WHITE, PP_ALIGN.CENTER)
        body = round_box(s, x, Inches(2.75), Inches(4.05), Inches(3.15), WHITE, LINE)
        multi(s, x + Inches(0.25), Inches(3.0), Inches(3.55), Inches(2.7), ["• " + b for b in bullets], 14, INK2, space=10)
    tb(s, Inches(0.45), Inches(6.1), Inches(12.4), Inches(0.85),
       "Flywheel: a university or institute that standardizes on Recrulyn becomes a talent channel; every company that hires from them wants the same inbox-to-offer path.",
       14, False, INK2)
    footer(s, 13, total)

    # ── 14 Why us ────────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    bg(s)
    header_bar(s, "Why Recrulyn wins this pitch", "Built around how HR actually works — starting with mail")
    points = [
        ("Mail-native, not portal-first", "Candidates and TPOs already email CVs. We do not ask the market to change behavior on day one."),
        ("AI where it pays", "Parsing, JD match, and recruiter ranking — not a chatbot that still needs a spreadsheet."),
        ("Paperwork is product", "LOA, NDA, LOC, LOR, certificates and approvals are first-class, not an afterthought PDF."),
        ("Role-correct access", "HR acts. Employees self-serve leave. Management sees pipeline and risk. Admin governs."),
        ("Fits Indian hiring reality", "Intern batches, campus drives, WhatsApp-forwarded CVs that still hit HR mail, multi-letter onboarding."),
        ("Deployable as a workspace", "HR inbox, talent, AI, documents, leave, analytics — one login, one story for buyers."),
    ]
    for i, (title, desc) in enumerate(points):
        col, row = i % 2, i // 2
        x, y = Inches(0.4) + Inches(col * 6.45), Inches(1.5) + Inches(row * 1.75)
        c = round_box(s, x, y, Inches(6.25), Inches(1.58), WHITE, LINE)
        tb(s, x + Inches(0.25), y + Inches(0.18), Inches(5.8), Inches(0.4), title, 16, True, INK)
        tb(s, x + Inches(0.25), y + Inches(0.62), Inches(5.8), Inches(0.8), desc, 13, False, INK2)
    footer(s, 14, total)

    # ── 15 Ask ───────────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    bg(s)
    header_bar(s, "The ask", "Three partnership shapes — pick the door that matches the room")
    asks = [
        ("Campus pilot (8–12 weeks)", [
            "One university drive or intern season",
            "HR mailbox connected; 1–2 JDs live",
            "Success: time-to-shortlist and letter turnaround",
            "Then: annual campus license + company seats",
        ]),
        ("Institute offtake", [
            "One graduating / trained cohort",
            "Joint shortlist with 1–2 employers",
            "Success: % of cohort with ranked, letter-ready files",
            "Then: recurring season + employer network",
        ]),
        ("Industry deploy", [
            "Connect IMAP in the first week",
            "Run AI recruiter on the next open role",
            "Success: hours saved per 25 applicants",
            "Then: documents + leave + management seats",
        ]),
    ]
    for i, (title, bullets) in enumerate(asks):
        x = Inches(0.4) + Inches(i * 4.25)
        h = box(s, x, Inches(1.5), Inches(4.05), Inches(0.7), ORANGE if i == 1 else NAVY)
        tb(s, x, Inches(1.6), Inches(4.05), Inches(0.5), title, 15, True, WHITE, PP_ALIGN.CENTER)
        c = round_box(s, x, Inches(2.3), Inches(4.05), Inches(3.55), WHITE, LINE)
        multi(s, x + Inches(0.22), Inches(2.5), Inches(3.6), Inches(3.2), ["• " + b for b in bullets], 14, INK2, space=10)
    tb(s, Inches(0.45), Inches(6.05), Inches(12.4), Inches(0.9),
       "What we need from you in this meeting: a named mailbox (or alias), one live JD, and a date for a 30-applicant dry run. Recrulyn brings the workspace, IMAP engine, and letter packs.",
       14, True, INK)
    footer(s, 15, total)

    # ── 16 Close ─────────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    bg(s, NAVY)
    box(s, 0, 0, Inches(0.18), H, ORANGE)
    add_logo(s, Inches(0.55), Inches(0.5), Inches(0.85))
    tb(s, Inches(1.55), Inches(0.68), Inches(10), Inches(0.5), "RECRULYN", 20, True, WHITE)
    tb(s, Inches(0.55), Inches(1.7), Inches(12), Inches(1.5),
       "Stop managing hiring in the inbox.\nStart running HR from the inbox.",
       28, True, WHITE)
    tb(s, Inches(0.55), Inches(3.45), Inches(12), Inches(1.0),
       "IMAP captures every CV. AI ranks the fit. HR acts from one workspace.\nUniversities, institutes, and industry finally share the same hiring operating system.",
       16, False, RGBColor(0xCB, 0xD5, 0xE1))

    close_pills = [
        "Connect the HR mailbox",
        "Run one JD",
        "Issue the first LOA from the same profile",
    ]
    for i, label in enumerate(close_pills):
        x = Inches(0.55) + Inches(i * 4.1)
        round_box(s, x, Inches(4.85), Inches(3.85), Inches(0.7), ORANGE)
        tb(s, x, Inches(4.98), Inches(3.85), Inches(0.5), f"{i+1}.  {label}", 14, True, WHITE, PP_ALIGN.CENTER)

    tb(s, Inches(0.55), Inches(6.2), Inches(12), Inches(0.7),
       "Let’s pick a campus drive, a training cohort, or an open industry role — and make the next 25 applicants effortless.",
       15, False, RGBColor(0x9C, 0xA3, 0xAF))
    box(s, 0, Inches(7.28), W, Inches(0.22), ORANGE)

    prs.save(str(OUT))
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
